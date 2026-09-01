#!/usr/bin/env python3
"""Собрать входы для сверки с эталонами.

В проекте два размеченных набора, и они меряют разное — поэтому оба нужны.

``AVTOZAP_ETALON_GOTOVYY.csv`` (469 заявок) — основной набор точности. У каждой
строки есть правильный код; заявок, где верный ответ «кода быть не должно», в
нём нет вовсе.

``AVTOZAP_etalon_200.xlsx`` (200 заявок) — набор поведения. Только в нём есть
27 заявок, где правильный ответ — отказ или переспрос, и только в нём записано,
где была права боевая система (128 из 200). Без него правило «не уверен —
переспроси, не знаешь названия — проси фото» проверить не на чем.

Из обоих берётся ТОЛЬКО сырой текст покупателя. Перевод, названия деталей и
ответы боевой системы идут в справочные поля и на вход конвейера не подаются:
перевод во время работы неоткуда взять, а подсказывать конвейеру ответ — значит
мерить не его.

    python scripts/build_etalon_input.py
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REFERENCE = os.path.join(ROOT, "data", "reference")

CSV_SRC = os.path.join(REFERENCE, "AVTOZAP_ETALON_GOTOVYY.csv")
CSV_OUT = os.path.join(ROOT, "data", "etalon_469.jsonl")
XLSX_SRC = os.path.join(REFERENCE, "AVTOZAP_etalon_200.xlsx")
XLSX_OUT = os.path.join(ROOT, "data", "etalon_200.jsonl")

#: Ответ «кода быть не должно» кодируется этим значением.
NO_CODE = "UNKNOWN"


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def build_from_csv(src: str, out: str, duplicates: dict[str, str],
                   known_codes: set[str]) -> dict:
    """Основной набор точности: 469 заявок, у каждой есть правильный код."""
    with open(src, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))

    records: list[dict] = []
    unknown_codes: dict[str, int] = {}
    canonicalised = 0
    for index, row in enumerate(rows):
        text = _text(row.get("raw_text"))
        if not text:
            continue
        expected = _text(row.get("expected_external_code"))
        if expected in duplicates:
            expected = duplicates[expected]
            canonicalised += 1
        # Код, которого нет в нашей версии словаря, помечаем — такую строку
        # нельзя ни засчитать, ни провалить, и молча ронять её нельзя тоже.
        scorable = bool(expected) and expected in known_codes
        if expected and not scorable:
            unknown_codes[expected] = unknown_codes.get(expected, 0) + 1
        records.append({
            "rfq_id": f"etalon469-{index + 1:03d}",
            "original_text": text,
            "expected_external_code": expected or NO_CODE,
            "scorable": scorable,
            # Справочно — на вход конвейера НЕ подаётся.
            "reference_translation": _text(row.get("translation")),
            "reference_name_ru": _text(row.get("name_ru")),
            "reference_name_az": _text(row.get("name_az")),
            "reference_label_source": _text(row.get("source")),
            "source_id": _text(row.get("id")),
        })

    _write(out, records)
    return {"rows": len(records), "canonicalised": canonicalised,
            "unknown_codes": unknown_codes,
            "unscorable": sum(1 for r in records if not r["scorable"])}


def build_from_xlsx(src: str, out: str, duplicates: dict[str, str]) -> dict:
    """Набор поведения: 200 заявок, из них 27 с ответом «кода быть не должно»."""
    import openpyxl

    workbook = openpyxl.load_workbook(src, read_only=True, data_only=True)
    rows = list(workbook[workbook.sheetnames[0]].iter_rows(values_only=True))
    header = [_text(h) for h in rows[0]]

    records: list[dict] = []
    for index, raw in enumerate(rows[1:]):
        if not any(cell is not None for cell in raw):
            continue
        row = dict(zip(header, raw))
        text = _text(row.get("текст_заявки"))
        if not text:
            continue
        expected = _text(row.get("ПРАВИЛЬНЫЙ_КОД"))
        expected = duplicates.get(expected, expected)
        records.append({
            "rfq_id": f"etalon-{index + 1:03d}",
            "original_text": text,
            "expected_external_code": expected or NO_CODE,
            "scorable": True,
            "reference_production_code": _text(row.get("что_ответила_система")),
            "reference_production_status": _text(row.get("статус")),
            "reference_production_correct": _text(row.get("система_права")),
            "reference_label_confidence": _text(row.get("уверенность")),
            "reference_comment": _text(row.get("комментарий")),
            "source_rfq_id": _text(row.get("rfq_id")),
        })

    _write(out, records)
    return {
        "rows": len(records),
        "no_code": sum(1 for r in records
                       if r["expected_external_code"] == NO_CODE),
        "production_baseline": sum(
            1 for r in records
            if r.get("reference_production_correct", "").lower() == "да"),
    }


def _write(path: str, records: list[dict]) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        for record in records:
            fh.write(json.dumps(record, ensure_ascii=False) + "\n")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--csv-src", default=CSV_SRC)
    parser.add_argument("--csv-out", default=CSV_OUT)
    parser.add_argument("--xlsx-src", default=XLSX_SRC)
    parser.add_argument("--xlsx-out", default=XLSX_OUT)
    args = parser.parse_args(argv)

    sys.path.insert(0, ROOT)
    from avtozap.dictionary import load, load_duplicates

    duplicates = load_duplicates()
    known = set(load().parts)

    if os.path.exists(args.csv_src):
        stats = build_from_csv(args.csv_src, args.csv_out, duplicates, known)
        print(f"{args.csv_out}: {stats['rows']} заявок (набор точности)")
        if stats["canonicalised"]:
            print(f"  кодов-дублей приведено к каноническим: "
                  f"{stats['canonicalised']}")
        if stats["unknown_codes"]:
            print(f"  ⚠ строк с кодом вне словаря 571: {stats['unscorable']}")
            for code, count in sorted(stats["unknown_codes"].items()):
                print(f"      {code} — {count} строк(и); в подсчёт не идут")

    if os.path.exists(args.xlsx_src):
        stats = build_from_xlsx(args.xlsx_src, args.xlsx_out, duplicates)
        print(f"{args.xlsx_out}: {stats['rows']} заявок (набор поведения)")
        print(f"  с ответом «кода быть не должно»: {stats['no_code']}")
        print(f"  точка отсчёта, боевая система  : {stats['production_baseline']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
