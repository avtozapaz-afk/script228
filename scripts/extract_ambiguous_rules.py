#!/usr/bin/env python3
"""Вынуть лист ``AMBIGUOUS_RULES`` словаря в ``data/ambiguous_rules.json``.

Правила пишет проект, а не я: в словаре v7 появился лист, где четыре термина
объявлены неоднозначными, к каждому указана пара кодов, правило для матчера и
готовый уточняющий вопрос покупателю. Файл собирается скриптом, а не руками,
чтобы при следующей поставке словаря его можно было пересобрать одной командой
и ничего не потерять и не переписать по памяти.

    python scripts/extract_ambiguous_rules.py
"""

from __future__ import annotations

import argparse
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "data", "AVTOZAP_slovar_FINAL_571.xlsx")
OUT = os.path.join(ROOT, "data", "ambiguous_rules.json")
SHEET = "AMBIGUOUS_RULES"

#: Термин записан вариантами через косую черту: «sveler / sveller / şveller».
VARIANT_SEPARATOR = "/"


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def extract(src: str) -> list[dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        return []
    rows = list(workbook[SHEET].iter_rows(values_only=True))
    header = [_text(h) for h in rows[0]]
    index = {name: position for position, name in enumerate(header)}

    def cell(row, name: str) -> str:
        position = index.get(name)
        return _text(row[position]) if position is not None else ""

    rules: list[dict] = []
    for row in rows[1:]:
        if not any(cell is not None for cell in row):
            continue
        if cell(row, "Статус").upper() != "AMBIGUOUS":
            continue
        variants = [part.strip()
                    for part in cell(row, "Термин/варианты").split(VARIANT_SEPARATOR)
                    if part.strip()]
        if not variants:
            continue
        rules.append({
            "variants": variants,
            "codes": [c for c in (cell(row, "Код 1"), cell(row, "Код 2")) if c],
            "parts": [p for p in (cell(row, "Деталь 1"), cell(row, "Деталь 2")) if p],
            "matcher_rule": cell(row, "Правило matcher"),
            "question": cell(row, "Уточняющий вопрос"),
        })
    return rules


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--src", default=SRC)
    parser.add_argument("--out", default=OUT)
    args = parser.parse_args(argv)

    rules = extract(args.src)
    payload = {
        "_note": ("Неоднозначные термины словаря. Собрано скриптом "
                  "scripts/extract_ambiguous_rules.py из листа AMBIGUOUS_RULES "
                  "файла словаря — руками не редактировать, пересобрать заново."),
        "rules": rules,
    }
    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=1)
        fh.write("\n")

    print(f"{args.out}: правил {len(rules)}")
    for rule in rules:
        print(f"  {' / '.join(rule['variants'])} -> "
              f"{' или '.join(rule['codes'])}: {rule['question']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
