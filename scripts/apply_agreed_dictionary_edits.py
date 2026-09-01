#!/usr/bin/env python3
"""Согласованные правки словаря, вносимые поверх поставки.

Поставка словаря — источник истины, и правится она **только** здесь, скриптом,
чтобы правка была воспроизводима и видна. Оригинал поставки лежит рядом как
``data/reference/AVTOZAP_slovar_v8_FINAL_original.xlsx``: в любой момент можно
сравнить и убедиться, что изменено ровно то, что перечислено ниже.

Правка одна, согласована 01.09.

**Убрать форму ``aktivatur turbnun ustunde olan`` у MU-038.** Это не термин, а
целое предложение, и вместе с ним в индекс попадали служебные слова ``üstündə``
и ``olan``. После этого они начинали «матчить» любую заявку с этими словами:
KZ-083 «Лючок топливного бака» упал с 0.9999 до 0.88 и пропустил вперёд EL-021,
хотя сам в словаре не менялся вовсе. Замена — контекстное правило
``aktivatur + turb`` → MU-038 в ``data/context_rules_local.json``: оно проверяет
сочетание слов и в индекс ничего не добавляет.

    python scripts/apply_agreed_dictionary_edits.py [--check]

``--check`` ничего не пишет, только сообщает, применены правки или нет.
"""

from __future__ import annotations

import argparse
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DICT = os.path.join(ROOT, "data", "AVTOZAP_slovar_FINAL_571.xlsx")

#: (код детали, точная форма) — строки, которые надо убрать из terms_normalized.
REMOVE_TERMS = (
    ("MU-038", "aktivatur turbnun ustunde olan"),
)


def _synonym_forms(value: str, form: str) -> str:
    """Убрать форму из строки синонимов, сохранив разделители."""
    parts = [p.strip() for p in str(value or "").split(";")]
    kept = [p for p in parts if p and p.strip().lower() != form.lower()]
    return "; ".join(kept)


def apply(path: str = DICT, check_only: bool = False) -> dict:
    import openpyxl

    workbook = openpyxl.load_workbook(path)
    terms = workbook["terms_normalized"]
    header = [cell.value for cell in terms[1]]
    code_at = header.index("external_code") + 1
    term_at = header.index("term") + 1

    doomed: list[int] = []
    for row in range(2, terms.max_row + 1):
        pair = (terms.cell(row, code_at).value, terms.cell(row, term_at).value)
        if pair in REMOVE_TERMS:
            doomed.append(row)

    parts = workbook["parts_synonyms"]
    parts_header = [cell.value for cell in parts[1]]
    parts_code_at = parts_header.index("external_code") + 1
    synonyms_at = parts_header.index("synonyms") + 1
    synonym_edits: list[int] = []
    for row in range(2, parts.max_row + 1):
        code = parts.cell(row, parts_code_at).value
        for target, form in REMOVE_TERMS:
            if code != target:
                continue
            cell = parts.cell(row, synonyms_at)
            trimmed = _synonym_forms(cell.value, form)
            if trimmed != (cell.value or ""):
                synonym_edits.append(row)
                if not check_only:
                    cell.value = trimmed

    if not check_only and (doomed or synonym_edits):
        for row in reversed(doomed):
            terms.delete_rows(row)
        workbook.save(path)

    return {"terms_removed": len(doomed), "synonyms_trimmed": len(synonym_edits)}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dict", default=DICT)
    parser.add_argument("--check", action="store_true",
                        help="только проверить, ничего не менять")
    args = parser.parse_args(argv)

    stats = apply(args.dict, check_only=args.check)
    if args.check:
        if stats["terms_removed"] or stats["synonyms_trimmed"]:
            print("правки ещё НЕ применены: "
                  f"терминов к удалению {stats['terms_removed']}, "
                  f"синонимов к правке {stats['synonyms_trimmed']}")
            return 1
        print("правки применены — убирать нечего")
        return 0
    print(f"удалено терминов: {stats['terms_removed']}, "
          f"поправлено строк синонимов: {stats['synonyms_trimmed']}")
    for code, form in REMOVE_TERMS:
        print(f"  {code}: убрана форма {form!r}")
    print("Замена работает контекстным правилом из data/context_rules_local.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
